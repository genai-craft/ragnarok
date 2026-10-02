"""ragnarok のエンジン: 文書の索引 (ページ本文 + 埋め込み + レイアウト木) と、質問への 3 つの動き。
  retrieve : ベクトル top-N → 確率判定を 2 段 (25 ずつ → 10 → k)。各段の確率と「どれでもない」を返す (棄却に使う)
  sweep    : 話題に触れているページを全部 (全ページ yes/no 判定、網羅モード = GraphRAG のグローバル検索の代わり)
  answer   : 上位ページを LLM に渡して回答 (思考あり/なし)
判定は openvons の LLMBackend (vLLM の logprob)。"""
from __future__ import annotations

import asyncio
import json
import os
import re
import time
from dataclasses import dataclass, field, asdict

import httpx
import numpy as np

from openvons.core.primitives import Option, Question
from openvons.lm.backends.llm_backend import LLMBackend

from .tree import Doc
from .layout import tree_from_layout


@dataclass
class Stage:
    name: str
    candidates: list[int]          # ページ (0 始まり)
    probs: list[float]             # 候補ごとの確率 (none を除いた相対)
    none: float                    # どれでもない
    latency_s: float


@dataclass
class Retrieval:
    pages: list[int]
    stages: list[Stage]
    emb_scores: dict[int, float]
    none_final: float
    abstain: bool
    section_path: dict[int, str] = field(default_factory=dict)

    def to_dict(self):
        d = asdict(self); d["emb_scores"] = {str(k): v for k, v in self.emb_scores.items()}; d["section_path"] = {str(k): v for k, v in self.section_path.items()}; return d


class Index:
    """1 文書の索引。pages / emb / tree。ディレクトリに保存・読込できる。"""
    def __init__(self, name: str, pages: list[str], emb: np.ndarray, tree: Doc | None = None, meta: dict | None = None):
        self.name, self.pages, self.emb, self.tree, self.meta = name, pages, emb, tree, meta or {}

    def save(self, d: str):
        os.makedirs(d, exist_ok=True)
        json.dump({"name": self.name, "pages": self.pages, "meta": self.meta}, open(f"{d}/pages.json", "w"), ensure_ascii=False)
        np.save(f"{d}/emb.npy", self.emb.astype(np.float32))
        if self.tree is not None:
            t = Doc(self.tree.name, [], self.tree.root); t.save(f"{d}/tree.json")

    @staticmethod
    def load(d: str) -> "Index":
        j = json.load(open(f"{d}/pages.json")); tree = None
        if os.path.exists(f"{d}/tree.json"):
            tree = Doc.load(f"{d}/tree.json"); tree.pages = j["pages"]
        return Index(j["name"], j["pages"], np.load(f"{d}/emb.npy"), tree, j.get("meta"))

    def section_path(self, p: int) -> str:
        if self.tree is None: return ""
        best = []
        def walk(n, path):
            nonlocal best
            if n.id != "root" and n.start - 1 <= p <= n.end - 1 and len(path) + 1 > len(best): best = path + [n.title]
            for c in n.children: walk(c, path + ([n.title] if n.id != "root" else []))
        walk(self.tree.root, [])
        return " > ".join(best)


class Engine:
    def __init__(self, embed_model: str = "Qwen/Qwen3-Embedding-0.6B", judge_url: str = "http://127.0.0.1:8310/v1", judge_model: str = "qwen27b",
                 fast_url: str = "http://127.0.0.1:8311/v1", fast_model: str = "qwen3-4b", answer_url: str | None = None, answer_model: str | None = None, device: str = "cuda"):
        from sentence_transformers import SentenceTransformer
        self.emb = SentenceTransformer(embed_model, device=device); self.emb.max_seq_length = 2048
        self.judge = LLMBackend(judge_url, judge_model, mode="logprob", concurrency=16)
        self.fast = LLMBackend(fast_url, fast_model, mode="logprob", concurrency=64)
        self.answer_url = (answer_url or judge_url).rstrip("/"); self.answer_model = answer_model or judge_model
        self.http = httpx.AsyncClient(timeout=600)

    # ---------------- 索引 ----------------
    def index_pdf(self, pdf: str, name: str | None = None, tree: bool = True) -> Index:
        import pymupdf
        name = name or os.path.basename(pdf).rsplit(".", 1)[0]
        with pymupdf.open(pdf) as d:
            pages = [p.get_text() for p in d]
        emb = self.emb.encode([p[:6000] or " " for p in pages], batch_size=8, normalize_embeddings=True, show_progress_bar=False)
        t = None
        if tree:
            try:
                t = tree_from_layout(pdf, name)
            except Exception:
                t = None
        return Index(name, pages, emb.astype(np.float32), t, {"source": pdf, "pages": len(pages)})

    @staticmethod
    def read_any(path: str) -> tuple[list[str], str]:
        """PDF 以外も「ページ」の列にする: DOCX (見出し/約 2,000 字で区切る)、PPTX (スライド = ページ)、HTML/TXT/MD (見出しか約 2,000 字)。戻り値 (pages, kind)。"""
        ext = os.path.splitext(path)[1].lower()
        if ext == ".pdf":
            import pymupdf
            with pymupdf.open(path) as d:
                return [p.get_text() for p in d], "pdf"
        if ext == ".pptx":
            from pptx import Presentation
            out = []
            for i, sl in enumerate(Presentation(path).slides):
                txt = "\n".join(sh.text_frame.text for sh in sl.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip())
                notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide and sl.notes_slide.notes_text_frame else ""
                out.append(f"[slide {i+1}]\n{txt}\n{('Notes: ' + notes) if notes.strip() else ''}")
            return out, "pptx"
        if ext == ".docx":
            import docx
            d = docx.Document(path); blocks = []
            for para in d.paragraphs:
                t = para.text.strip()
                if not t: continue
                blocks.append(("h" if para.style.name.lower().startswith("heading") else "p", t))
            for tb in d.tables:
                blocks.append(("p", "\n".join(" | ".join(c.text.strip() for c in row.cells) for row in tb.rows)))
            return Engine._chunk_blocks(blocks), "docx"
        if ext in (".html", ".htm"):
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(open(path, encoding="utf-8", errors="ignore").read(), "html.parser")
            for t in soup(["script", "style", "nav", "footer"]): t.decompose()
            blocks = [("h" if el.name in ("h1", "h2", "h3") else "p", el.get_text(" ", strip=True)) for el in soup.find_all(["h1", "h2", "h3", "p", "li", "td", "pre"]) if el.get_text(strip=True)]
            return Engine._chunk_blocks(blocks), "html"
        # txt / md
        txt = open(path, encoding="utf-8", errors="ignore").read()
        blocks = [("h" if re.match(r"^#{1,3}\s", ln) else "p", ln.strip()) for ln in txt.splitlines() if ln.strip()]
        return Engine._chunk_blocks(blocks), ext.lstrip(".") or "text"

    @staticmethod
    def _chunk_blocks(blocks: list[tuple[str, str]], target: int = 2000) -> list[str]:
        """見出しで区切りつつ、長すぎる節は target 文字で割る。"""
        pages, cur = [], ""
        for kind, t in blocks:
            if (kind == "h" and len(cur) > 400) or len(cur) + len(t) > target * 1.5:
                pages.append(cur.strip()); cur = ""
            cur += ("\n\n" if cur else "") + t
            while len(cur) > target * 1.5:
                pages.append(cur[:target].strip()); cur = cur[target:]
        if cur.strip(): pages.append(cur.strip())
        return pages or [""]

    def index_file(self, path: str, name: str | None = None) -> Index:
        """PDF / DOCX / PPTX / HTML / TXT / MD。PDF 以外はレイアウト木なし (見出しで区切った「ページ」)。"""
        if path.lower().endswith(".pdf"):
            return self.index_pdf(path, name)
        pages, kind = self.read_any(path)
        ix = self.index_pages(name or os.path.basename(path).rsplit(".", 1)[0], pages)
        ix.meta.update({"source": path, "format": kind, "unit": "slide" if kind == "pptx" else "section"}); return ix

    def index_pages(self, name: str, pages: list[str]) -> Index:
        emb = self.emb.encode([p[:6000] or " " for p in pages], batch_size=8, normalize_embeddings=True, show_progress_bar=False)
        return Index(name, pages, emb.astype(np.float32), None, {"pages": len(pages)})

    # ---------------- 検索 ----------------
    @staticmethod
    def _snip(t: str, n: int = 300) -> str:
        return re.sub(r"\s+", " ", t)[:n]

    @staticmethod
    def _grams(s: str) -> set[str]:
        s = re.sub(r"\s+", " ", s.lower())
        words = set(re.findall(r"[a-z0-9][a-z0-9.%$-]{1,}", s))
        cjk = re.sub(r"[^ぁ-んァ-ン一-龥ー]", "", s)
        return words | {cjk[i:i+2] for i in range(len(cjk) - 1)}

    @classmethod
    def snip_for(cls, q: str, text: str, n: int = 360) -> str:
        """質問に最も関係する窓を抜く (ページの先頭 300 字は見出しや定型文のことが多く、判定役が本文を見られない問題への対処)。
        英語は単語、日本語は 2 文字の重なりで窓を採点。先頭の見出し行 (80 字) は常に添える。"""
        t = re.sub(r"\s+", " ", text).strip()
        if len(t) <= n:
            return t
        qg = cls._grams(q)
        if not qg:
            return t[:n]
        step = max(60, n // 3); best, best_sc = 0, -1.0
        for i in range(0, len(t) - n // 2, step):
            w = t[i : i + n]; sc = len(qg & cls._grams(w))
            if sc > best_sc:
                best, best_sc = i, sc
        head = t[:80]
        body = t[best : best + n]
        return body if best == 0 else f"{head} … {body}"

    def _state(self, ix: Index) -> str:
        kind = ix.meta.get("kind") or "document"
        return f"Document: {ix.name} ({kind})"

    async def _decide(self, ix: Index, q: str, cand: list[int], backend: LLMBackend, with_path: bool) -> Stage:
        t0 = time.perf_counter()
        opts = [Option(f"p{p}", (f"[{ix.section_path(p)}] " if with_path and ix.tree else "") + f"page {p+1}: {self.snip_for(q, ix.pages[p])}") for p in cand]
        opts.append(Option("none", "None of these pages contains the answer"))
        qq = Question("choice", f"Which page contains the information needed to answer this question?\nQuestion: {q}", opts)
        d = (await backend.adecide(self._state(ix), [qq]))[0]
        ps = [float(x) for x in d.probs[:-1]]; s = sum(ps) or 1.0
        return Stage("", cand, [p / s for p in ps], float(d.probs[-1]), time.perf_counter() - t0)

    async def retrieve(self, ix: Index, q: str, k: int = 5, topn: int = 50, group: int = 25, pool_each: int = 5, fast_first: bool = False, with_path: bool = False,
                       abstain_th: float = 0.5) -> Retrieval:
        qe = self.emb.encode([q], prompt_name="query", normalize_embeddings=True, show_progress_bar=False)[0]
        sims = ix.emb @ qe; order = [int(i) for i in np.argsort(-sims)[:topn]]
        stages: list[Stage] = []; pool: list[int] = []
        if len(order) > group:
            b = self.fast if fast_first else self.judge
            res = await asyncio.gather(*[self._decide(ix, q, order[i : i + group], b, with_path) for i in range(0, len(order), group)])
            for j, st in enumerate(res):
                st.name = f"stage1-{j+1}"; stages.append(st)
                pool += [p for p, _ in sorted(zip(st.candidates, st.probs), key=lambda x: -x[1])[:pool_each]]
        else:
            pool = order
        fin = await self._decide(ix, q, pool, self.judge, with_path); fin.name = "final"; stages.append(fin)
        ranked = [p for p, _ in sorted(zip(fin.candidates, fin.probs), key=lambda x: -x[1])]
        pages = ranked[:k]
        return Retrieval(pages, stages, {int(p): float(sims[p]) for p in order}, fin.none, fin.none >= abstain_th, {p: ix.section_path(p) for p in pages})

    # ---------------- 網羅 (集約) ----------------
    async def sweep(self, ix: Index, topic: str, pages: list[int] | None = None, use_fast: bool = True, th: float = 0.5) -> dict[int, float]:
        b = self.fast if use_fast else self.judge
        idx = pages if pages is not None else list(range(len(ix.pages)))
        async def one(p):
            qq = Question("choice", f"Does this page substantively discuss {topic} (not just a passing mention)?", [Option("yes", "yes"), Option("no", "no")])
            d = (await b.adecide(f"{self._state(ix)}. Page {p+1}:\n{re.sub(r'[ \t]+', ' ', ix.pages[p])[:3500]}", [qq]))[0]
            return p, float(d.probs[0])
        return dict(await asyncio.gather(*[one(p) for p in idx]))

    # ---------------- 回答 ----------------
    async def chat(self, prompt: str, think: bool = False, max_tokens: int = 400) -> str:
        r = await self.http.post(f"{self.answer_url}/chat/completions", json={"model": self.answer_model, "messages": [{"role": "user", "content": prompt}],
                                 "max_tokens": 3000 if think else max_tokens, "temperature": 0, "chat_template_kwargs": {"enable_thinking": think}})
        m = r.json()["choices"][0]["message"]; txt = (m.get("content") or "").strip()
        return txt.split("</think>")[-1].strip() if "</think>" in txt else txt

    def _answer_prompt(self, ix: Index, q: str, pages: list[int], lang: str = "auto") -> str:
        ctx = "\n\n".join(f"=== page {p+1} ===\n{ix.pages[p][:6000]}" for p in pages)
        ja = lang == "ja" or (lang == "auto" and re.search(r"[ぁ-んァ-ン一-龥]", q))
        inst = ("以下は文書の抜粋です。抜粋だけを根拠に質問に答えてください。数値には単位を付け、必要なら計算過程を短く示し、根拠のページ番号を添えてください。抜粋に答えが無ければ「文書中に見つかりません」と答えてください。"
                if ja else "Using only the excerpts below, answer the question. Give numbers with units, show a brief calculation if needed, and cite page numbers. If the excerpts do not contain the answer, say \"Not found in the document\".")
        return f"{inst}\n\nDocument: {ix.name}\n\n{ctx}\n\nQuestion: {q}\nAnswer:"

    async def answer(self, ix: Index, q: str, pages: list[int], think: bool = True, lang: str = "auto") -> str:
        return await self.chat(self._answer_prompt(ix, q, pages, lang), think=think)

    async def answer_stream(self, ix: Index, q: str, pages: list[int], think: bool = True, lang: str = "auto"):
        """回答をトークン単位で流す。思考中は ("think", 文字数) を、本文は ("text", 断片) を yield。"""
        body = {"model": self.answer_model, "messages": [{"role": "user", "content": self._answer_prompt(ix, q, pages, lang)}], "max_tokens": 3000 if think else 500,
                "temperature": 0, "stream": True, "chat_template_kwargs": {"enable_thinking": think}}
        in_think = think; buf = ""
        async with self.http.stream("POST", f"{self.answer_url}/chat/completions", json=body) as r:
            async for line in r.aiter_lines():
                if not line.startswith("data: ") or line.strip() == "data: [DONE]": continue
                try:
                    d = json.loads(line[6:]); delta = d["choices"][0].get("delta", {})
                except Exception:
                    continue
                rc = delta.get("reasoning_content") or delta.get("reasoning")
                if rc:
                    yield ("think", len(rc)); continue
                c = delta.get("content") or ""
                if not c: continue
                if in_think:
                    # Qwen3 系は思考が content に流れてくる (<think> は雛形側で開かれ、本文は </think> の後)。</think> が来るまで思考として数える
                    buf += c
                    if "</think>" in buf:
                        in_think = False; after = buf.split("</think>", 1)[1].lstrip("\n"); buf = ""
                        if after: yield ("text", after)
                    else:
                        yield ("think", len(c))
                    continue
                yield ("text", c)
        if in_think and buf.strip():   # 思考モードなのに </think> が最後まで来なかった = 思考せずに答えた
            yield ("text", buf.replace("<think>", "").strip())

    async def summarize_hits(self, ix: Index, topic: str, hits: list[int], lang: str = "auto") -> str:
        ctx = "\n\n".join(f"=== page {p+1} ===\n{ix.pages[p][:2500]}" for p in hits[:20])
        ja = lang == "ja" or (lang == "auto" and re.search(r"[ぁ-んァ-ン一-龥]", topic))
        inst = f"以下は文書 {ix.name} のうち「{topic}」に触れているページの抜粋です。ページ番号を添えて箇条書きで要点をまとめてください。" if ja else f"Below are the pages of {ix.name} that discuss {topic}. Summarize the key points as bullets, citing page numbers."
        return await self.chat(f"{inst}\n\n{ctx}", think=False, max_tokens=700)
